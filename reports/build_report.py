"""Builds reports/Freight_Rate_Report.docx from the saved CV results and the scorer chart."""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
cv = pd.read_csv(ROOT / "outputs" / "cv_summary.csv")
meta = json.loads((ROOT / "outputs" / "validation_metrics.json").read_text())

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
doc.add_heading("Freight Rate Prediction - Approach & Validation Report", 0)

doc.add_heading("1. Data and problem", 1)
doc.add_paragraph(
    "Labeled data: 48,000 loads from 2025-01-01 to 2025-10-31. Target loads: 12,000 loads from "
    "2025-11-01 to 2025-12-31 (no labels). The target period is entirely in the FUTURE of the labeled data, "
    "so the model has to extrapolate in time, not only interpolate.")

doc.add_heading("2. Split and validation approach", 1)
for s in [
    "Random K-fold was rejected: it mixes months, lets the model see the market regime of the test period and gives optimistic scores.",
    "Rolling-origin temporal validation instead: train on months <= 6 and test on 7-8; train <= 7 test 8-9; train <= 8 test 9-10. Each fold mimics the real task (train on the past, predict the future).",
    "Corrupted-label filtering is fitted on the training part of each fold only (no test leakage). Errors are reported both on all rows and on non-corrupted labels, since ~1.4% of labels are unpredictable noise and dominate RMSE.",
    "Metrics: MAE, RMSE, MAPE and median APE. A baseline (median rate-per-mile by equipment x distance) is included.",
    "Unseen-city stress test: 8 cities never appear in train but appear in the target set (12% of target rows). I hold out 8 cities from training and score loads touching them.",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.add_heading("3. Data-quality issues and handling", 1)
rows = [
    ("Corrupted labels (677 rows, 1.4%)", "Rate per mile is 0.2-0.4x or 2.2-5.4x the expected value, spread evenly over months/equipment.", "Flagged with an out-of-fold robust model (|log residual| > 0.30) and removed from training only. They are not removed from evaluation of 'all rows'."),
    ("weight: 300 missing in train, 165 in target", "Missing values.", "Median by equipment (learned on train) + missing indicator."),
    ("weight: 292 negative values", "Sign errors (absolute value looks plausible).", "abs() + indicator."),
    ("weight: ~1,200 rows exactly at +/-47,500", "Values are capped/censored.", "Clipped to 47,500; tree model handles the pile-up."),
    ("market_index: 374 missing train / 249 target", "Missing; is a daily series plus small noise.", "Not used in the final model (see below)."),
    ("distance vs coordinates", "Reported distance is road distance (~1.18x haversine); 214 rows have ratio > 1.5.", "Added haversine and distance/haversine ratio features. No rows dropped."),
    ("Unseen cities", "8 cities in the target set are absent from train.", "No city-ID or target encoding; geography enters through coordinates (constant per city) and distance, so new cities work. Stress test shows only ~0.1 pp MAPE degradation."),
]
tbl = doc.add_table(rows=1, cols=3); tbl.style = "Light Grid Accent 1"
for c, h in zip(tbl.rows[0].cells, ["Issue", "What I found", "Handling"]):
    c.text = h
for r in rows:
    cells = tbl.add_row().cells
    for c, x in zip(cells, r):
        c.text = x
doc.add_paragraph()
doc.add_picture(str(ROOT / "reports" / "eda.png"), width=Inches(6.5))

doc.add_heading("4. Model", 1)
for s in [
    "Target: log(rate per mile). Rate is ~linear in distance (corr 0.91); modelling the per-mile rate removes that effect and makes errors relative. Predictions = exp(pred) x distance.",
    "Model: mean (in log space) of two scikit-learn HistGradientBoostingRegressor models, one with squared-error loss and one with absolute-error loss (robust to residual noise). Gradient boosting captures interactions such as equipment x distance x geography without manual engineering.",
    "Features: distance, log distance, cleaned weight (+ flags), equipment, pickup/delivery coordinates, haversine, distance ratio, day of week.",
    "Excluded on evidence: market_index (worse in every fold - it drifts and the target period sits at a level rarely seen in training) and quote_signal (see section 5).",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.add_heading("5. Results (mean over the 3 rolling folds)", 1)
show = cv[cv.scope == "clean_labels"].copy()
t2 = doc.add_table(rows=1, cols=5); t2.style = "Light Grid Accent 1"
for c, h in zip(t2.rows[0].cells, ["Configuration (clean labels)", "MAE $", "RMSE $", "MAPE", "MdAPE"]):
    c.text = h
for _, r in show.iterrows():
    cells = t2.add_row().cells
    cells[0].text = r["config"]; cells[1].text = f"{r.MAE:.1f}"; cells[2].text = f"{r.RMSE:.1f}"
    cells[3].text = f"{r.MAPE:.2%}"; cells[4].text = f"{r.MdAPE:.2%}"
doc.add_paragraph()
doc.add_picture(str(ROOT / "reports" / "cv_folds.png"), width=Inches(5.2))
doc.add_paragraph(
    "quote_signal judgement call: it helps strongly only in the last fold (MAPE 1.2% vs 1.8%), where the training window already "
    "contains the same regime, and hurts in the two older folds (3.4-4.0% vs 1.8-2.2%). Its daily mean switches between regimes "
    "(~1.9 / ~2.05 / ~2.25) and its relation to price is U-shaped, so it is not stable over time. Because the target period is unseen, "
    "I chose the option that is robust across all folds. Relative variants (signal minus daily mean, ratio) did not fix this.")
u = meta["unseen_city_test"]
doc.add_paragraph(
    f"Unseen-city test (last fold, clean labels): MAPE {u['cities seen in training']['MAPE']:.2%} with the cities in training vs "
    f"{u['cities NOT seen in training']['MAPE']:.2%} without them.")

doc.add_heading("6. Fixed December prediction chart", 1)
doc.add_paragraph(
    "Produced with the provided score.py. The chart inputs contain no market_index/quote_signal, which is one more reason the final "
    "model does not rely on them. Only the date changes, so the curve reflects the learned day-of-week pattern (~+/-1%). "
    "The model has no information about December-specific demand (holidays, seasonality) because no December labels exist; the chart should "
    "not be read as a forecast of trend.")
doc.add_picture(str(ROOT / "scorer_results" / "candidate_december.png"), width=Inches(6.5))

doc.add_heading("7. Limitations", 1)
for s in [
    "Label noise (1.4%) is unpredictable; all-row RMSE stays ~620 regardless of model.",
    "No labels for Nov-Dec: a market-level shift in these months cannot be learned and would add a bias to all predictions.",
    "The validation-set regime of quote_signal could make it useful; I kept it out for robustness, at a measurable cost in the best-case fold.",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.save(ROOT / "reports" / "Freight_Rate_Report.docx")
print("saved")
